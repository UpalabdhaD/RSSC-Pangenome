"""Build Supplementary Tables 7 and 8 for the dN/dS section.

Table 7 (supplement PDF, 13 rows): PAML M0 omega and HyPhy BUSTED results per lineage-aware class.
  Inputs: ../figure4_paml/omega_by_twilight_class.tsv
          ../figure4_paml/omega_wilcoxon_vs_collection_core.tsv
          busted_by_class.tsv
  Output: Table_S7_dnds_by_class.tsv (raw numbers; the supplement formats them)

Table 8 (Excel, 887 rows): families with M0 omega > 1 or BUSTED q <= 0.05. The iScience Final
Files Checklist asks for supplemental tables longer than 3 pages as Excel files.
  Inputs: selection_candidates.tsv, busted_function_groups.tsv, ../busted_gene_tree/busted_complete.tsv
  Output: Table_S8_selection_candidates.xlsx

Run from this folder: python3 -I make_supp_tables_7_8.py
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

FIG4 = Path("../figure4_paml")
TABLE7 = Path("Table_S7_dnds_by_class.tsv")
TABLE8 = Path("Table_S8_selection_candidates.xlsx")

# Class names as the manuscript writes them
CLASS_NAMES = {
    "Collection core": "Collection Core",
    "Absent in large lineages": "Absent in large lineages",
    "Core and intermediate": "Core and Intermediate",
    "Collection intermediate": "Collection Intermediate",
    "Multi-lineage core": "Multi-lineage Core",
    "Core, intermediate and rare": "Core, Intermediate and Rare",
    "Intermediate and rare": "Intermediate and Rare",
    "Multi-lineage intermediate": "Multi-lineage Intermediate",
    "Multi-lineage rare": "Multi-lineage Rare",
    "Core and rare": "Core and Rare",
    "Lineage specific rare": "Lineage-specific Rare",
    "Lineage specific core": "Lineage-specific Core",
    "Lineage specific intermediate": "Lineage-specific Intermediate",
}
TREE_NAMES = {"pruned_species_tree": "Pruned core-genome SNP phylogeny", "gene_tree": "FastTree gene tree"}


def read_tsv(path):
    with open(path, newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# ---- Table 7 ----------------------------------------------------------------------------------
omega = read_tsv(FIG4 / "omega_by_twilight_class.tsv")
wilcox = {r["specific_class"]: r for r in read_tsv(FIG4 / "omega_wilcoxon_vs_collection_core.tsv")}
busted = {r["specific_class"]: r for r in read_tsv("busted_by_class.tsv")}
assert len(omega) == 13 and set(CLASS_NAMES) == {r["specific_class"] for r in omega}
assert sum(int(r["n"]) for r in omega) == 9400

with open(TABLE7, "w", newline="") as out:
    writer = csv.writer(out, delimiter="\t", lineterminator="\n")
    writer.writerow(["class", "families_m0", "median_omega", "q1", "q3", "omega_above_1",
                     "pct_omega_above_1", "p_bh_vs_collection_core", "busted_tested",
                     "busted_positive", "pct_busted_positive"])
    for r in omega:                                    # already ordered by median omega
        cls, n = r["specific_class"], int(r["n"])
        above = round(float(r["percent_omega_above_1"]) * n / 100)
        b = busted[cls]
        writer.writerow([CLASS_NAMES[cls], n, r["median_omega"], r["q1"], r["q3"], above,
                         r["percent_omega_above_1"], wilcox[cls]["p_bh"] if cls in wilcox else "",
                         b["families"], b["positive"], b["pct"]])
print(f"Wrote {TABLE7}")

# ---- Table 8 ----------------------------------------------------------------------------------
candidates = read_tsv("selection_candidates.tsv")
assert len(candidates) == 887
groups = {r["gene"]: r["group"] for r in read_tsv("busted_function_groups.tsv") if r["positive"] == "TRUE"}
busted_all = {r["gene"]: r for r in read_tsv("../busted_gene_tree/busted_complete.tsv")}

columns = ["Gene family (Panaroo)", "Lineage-aware class", "Broad class", "Distinct sequences",
           "Tree", "M0 omega", "M0 status", "BUSTED P", "BUSTED q (BH)", "BUSTED global omega",
           "Evidence", "Function group", "COG category", "Preferred name (eggNOG)",
           "Description (eggNOG)"]
rows = []
for r in candidates:
    m0 = number(r["omega"])
    tree = r["tree_source"] if r["tree_source"] not in ("", "NA") else r["busted_tree"]
    seqs = r["sequence_count"] if r["sequence_count"] not in ("", "NA") else busted_all[r["gene"]]["n_sequences"]
    rows.append([
        r["gene"], CLASS_NAMES.get(r["specific_class"], r["specific_class"]), r["general_class"],
        int(seqs), TREE_NAMES[tree], m0,
        "Estimated" if m0 is not None else "Boundary value omega = 999; excluded",
        number(r["busted_pvalue"]), number(r["busted_qvalue"]), number(r["global_omega"]),
        {"BUSTED": "BUSTED q <= 0.05", "M0 omega > 1": "M0 omega > 1",
         "M0 omega > 1 and BUSTED": "Both"}[r["evidence"]],
        groups.get(r["gene"], ""),
        *("" if r[k] in ("NA", "-") else r[k] for k in ("COG_category", "Preferred_name", "Description")),
    ])
evidence_order = {"Both": 0, "M0 omega > 1": 1, "BUSTED q <= 0.05": 2}
rows.sort(key=lambda x: (evidence_order[x[10]], -(x[5] or 0) if x[10] != "BUSTED q <= 0.05" else (x[8] or 1)))

legend = [
    "Table S8. Gene families with PAML M0 omega > 1 or significant HyPhy BUSTED support, related to Figure 4.",
    "The table lists 887 gene families: 5 with both M0 omega > 1 and BUSTED q <= 0.05, 215 with M0 omega > 1 only, and 667 with BUSTED q <= 0.05 only. "
    "Rows are sorted by evidence, then by M0 omega (descending) or BUSTED q (ascending).",
    "Gene family: Panaroo gene-family name; ~~~ joins names that Panaroo merged into one family.",
    "Lineage-aware class: one of the 13 Twilight classes defined in STAR Methods. Broad class: Core, Intermediate or Rare; Varied marks families whose frequency class differs between lineages (for example Core and Intermediate).",
    "Distinct sequences: number of distinct coding sequences in the deduplicated codon alignment.",
    "Tree: tree used for codeml and BUSTED, either the core-genome SNP phylogeny pruned to the family or, for families in which a strain contributed two or more distinct sequences, a FastTree gene tree.",
    "M0 omega: family-wide dN/dS from the one-ratio model (PAML codeml). One BUSTED-positive family (group_14449) returned the boundary value omega = 999, which leaves omega undefined; its M0 omega is left blank.",
    "BUSTED P and BUSTED q: likelihood-ratio test P value and Benjamini-Hochberg q value across the 6,804 families with a BUSTED result. Blank: the family had fewer than four distinct sequences or BUSTED did not finish.",
    "BUSTED global omega: dN/dS estimated by HyPhy under the MG94xREV codon model before the branch-site fit.",
    "Function group: BUSTED-positive families assigned to the type III effector, Rhs/VgrG/toxin, or cyclic-di-GMP groups by the text rules in STAR Methods.",
    "COG category, Preferred name and Description: eggNOG-mapper v2.1.12 annotation; blank when eggNOG-mapper gave none.",
]

wb = Workbook()
ws = wb.active
ws.title = "Table S8"
ws.append(columns)
for c in ws[1]:
    c.font = Font(bold=True)
for row in rows:
    ws.append(row)
for col, width in zip("ABCDEFGHIJKLMNO", [22, 28, 12, 10, 30, 10, 34, 12, 12, 12, 18, 24, 10, 22, 60]):
    ws.column_dimensions[col].width = width
for col in "FHIJ":
    for cell in ws[col][1:]:
        cell.number_format = "0.000" if col in "FJ" else "0.00E+00"
ws.freeze_panes = "B2"
lg = wb.create_sheet("Legend")
for line in legend:
    lg.append([line])
lg["A1"].font = Font(bold=True)
lg.column_dimensions["A"].width = 150
for cell in lg["A"]:
    cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(TABLE8)
print(f"Wrote {TABLE8}: {len(rows)} rows")
